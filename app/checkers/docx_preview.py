"""Сборка структуры документа для показа в браузере.

Зачем это нужно. Отчёт по коду показывает исходник с подсветкой: видно,
на какой строке замечание. Для документа такого не было, и пользователь
читал «Абзац 17, красная строка 1,0 см», а потом искал этот абзац глазами
в Word. Модуль собирает представление документа, по которому фронтенд
строит страницу с подсветкой проблемных мест.

Это приближение, а не точная копия вёрстки Word: показываются текст,
выравнивание, начертание и уровень заголовка. Для поиска места этого
достаточно, а точный постраничный рендер потребовал бы LibreOffice на
сервере.

Индексы абзацев совпадают с нумерацией в `document.paragraphs`, поэтому
по ним замечание однозначно связывается со своим местом в документе.
"""
from __future__ import annotations

from typing import Any, Iterator

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# Предел на случай очень больших документов: дальше показывать смысла нет,
# страница в браузере станет неподъёмной.
MAX_BLOCKS = 2000
MAX_TABLE_ROWS = 20
MAX_CELL_CHARS = 200

_ALIGN_NAMES = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _iter_blocks(document) -> Iterator[Paragraph | Table]:
    """Перебирает абзацы и таблицы в том порядке, в каком они идут в документе.

    python-docx отдаёт `document.paragraphs` и `document.tables` раздельно,
    из-за чего таблица теряет своё место в тексте. Поэтому идём по дочерним
    элементам тела документа напрямую.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int:
    """Уровень заголовка: 1-3, либо 0 для обычного абзаца."""
    style_name = (paragraph.style.name or "").lower()
    if not ("heading" in style_name or "заголов" in style_name):
        return 0
    for digit in ("1", "2", "3"):
        if digit in style_name:
            return int(digit)
    return 1


def _is_bold(paragraph: Paragraph) -> bool:
    """Абзац считается полужирным, если таким набран весь непустой текст."""
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    return all(run.bold or (run.style is not None and "bold" in (run.style.name or "").lower())
               for run in runs)


def _paragraph_block(paragraph: Paragraph, index: int) -> dict[str, Any]:
    text = paragraph.text
    return {
        "kind": "p",
        "index": index,
        "text": text,
        "align": _ALIGN_NAMES.get(paragraph.alignment, "left"),
        "heading": _heading_level(paragraph),
        "bold": _is_bold(paragraph),
        "empty": not text.strip(),
    }


def _table_block(table: Table) -> dict[str, Any]:
    rows: list[list[str]] = []
    for row in table.rows[:MAX_TABLE_ROWS]:
        cells = []
        for cell in row.cells:
            value = cell.text.strip()
            if len(value) > MAX_CELL_CHARS:
                value = value[:MAX_CELL_CHARS] + "…"
            cells.append(value)
        rows.append(cells)
    return {
        "kind": "table",
        "rows": rows,
        "truncated_rows": max(0, len(table.rows) - MAX_TABLE_ROWS),
    }


def build_preview(document) -> dict[str, Any]:
    """Собирает представление документа: абзацы и таблицы по порядку.

    Returns:
        Словарь с полями `blocks` (список абзацев и таблиц) и `truncated`
        (сколько блоков не поместилось в предел).
    """
    blocks: list[dict[str, Any]] = []
    paragraph_index = -1
    overflow = 0

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            paragraph_index += 1
            if len(blocks) >= MAX_BLOCKS:
                overflow += 1
                continue
            blocks.append(_paragraph_block(block, paragraph_index))
        else:
            if len(blocks) >= MAX_BLOCKS:
                overflow += 1
                continue
            blocks.append(_table_block(block))

    return {"blocks": blocks, "truncated": overflow}
