"""Тесты предпросмотра документа.

Предпросмотр нужен для подсветки замечаний прямо в тексте: пользователь
видит, где именно нарушено оформление, вместо того чтобы искать «Абзац 17»
глазами в Word. Ключевое требование — индексы абзацев в предпросмотре
совпадают с полем para_index у замечаний. Если они разойдутся, подсветка
встанет не на тот абзац, и это хуже, чем её отсутствие.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.checkers.docx_checker import check_docx_document
from app.checkers.docx_preview import MAX_TABLE_ROWS, build_preview


def _paragraphs(preview) -> list[dict]:
    return [b for b in preview["blocks"] if b["kind"] == "p"]


class TestBuildPreview:
    def test_абзацы_в_порядке_документа(self):
        doc = Document()
        for text in ("Первый", "Второй", "Третий"):
            doc.add_paragraph(text)
        blocks = _paragraphs(build_preview(doc))
        assert [b["text"] for b in blocks] == ["Первый", "Второй", "Третий"]

    def test_индексы_совпадают_с_python_docx(self):
        doc = Document()
        for text in ("А", "", "Б", "В"):
            doc.add_paragraph(text)
        for block in _paragraphs(build_preview(doc)):
            assert doc.paragraphs[block["index"]].text == block["text"]

    def test_таблица_на_своём_месте_между_абзацами(self):
        doc = Document()
        doc.add_paragraph("До таблицы")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(1, 1).text = "B2"
        doc.add_paragraph("После таблицы")

        kinds = [b["kind"] for b in build_preview(doc)["blocks"]]
        assert kinds == ["p", "table", "p"]

    def test_индексы_не_сбиваются_из_за_таблицы(self):
        # Абзацы внутри таблицы не входят в document.paragraphs, поэтому
        # нумерация обязана остаться сплошной по абзацам тела документа
        doc = Document()
        doc.add_paragraph("До")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "ячейка"
        doc.add_paragraph("После")
        blocks = _paragraphs(build_preview(doc))
        assert [b["index"] for b in blocks] == [0, 1]
        assert doc.paragraphs[1].text == "После"

    def test_содержимое_таблицы(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Параметр"
        table.cell(0, 1).text = "Значение"
        block = build_preview(doc)["blocks"][0]
        assert block["rows"][0] == ["Параметр", "Значение"]
        assert block["truncated_rows"] == 0

    def test_длинная_таблица_обрезается(self):
        doc = Document()
        doc.add_table(rows=MAX_TABLE_ROWS + 5, cols=1)
        block = build_preview(doc)["blocks"][0]
        assert len(block["rows"]) == MAX_TABLE_ROWS
        assert block["truncated_rows"] == 5

    def test_выравнивание(self):
        doc = Document()
        doc.add_paragraph("обычный")
        doc.add_paragraph("по центру").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("по ширине").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        aligns = [b["align"] for b in _paragraphs(build_preview(doc))]
        assert aligns == ["left", "center", "justify"]

    def test_уровень_заголовка(self):
        doc = Document()
        doc.add_heading("Раздел", level=1)
        doc.add_heading("Подраздел", level=2)
        doc.add_paragraph("Текст")
        levels = [b["heading"] for b in _paragraphs(build_preview(doc))]
        assert levels == [1, 2, 0]

    def test_полужирный_абзац(self):
        doc = Document()
        bold = doc.add_paragraph()
        bold.add_run("весь текст жирный").bold = True
        mixed = doc.add_paragraph()
        mixed.add_run("жирное ").bold = True
        mixed.add_run("и обычное")
        blocks = _paragraphs(build_preview(doc))
        assert blocks[0]["bold"] is True
        assert blocks[1]["bold"] is False

    def test_пустой_абзац_помечен(self):
        doc = Document()
        doc.add_paragraph("текст")
        doc.add_paragraph("   ")
        blocks = _paragraphs(build_preview(doc))
        assert blocks[0]["empty"] is False
        assert blocks[1]["empty"] is True

    def test_пустой_документ(self):
        preview = build_preview(Document())
        assert preview["truncated"] == 0
        assert all(b["kind"] == "p" for b in preview["blocks"])


class TestPreviewInReport:
    def _doc_with_issues(self, tmp_path: Path) -> Path:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"      # не Times New Roman — будет замечание
        normal.font.size = Pt(11)
        doc.add_heading("Введение", level=1)
        for i in range(3):
            p = doc.add_paragraph(f"Абзац основного текста номер {i} с достаточной длиной.")
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
        path = tmp_path / "issues.docx"
        doc.save(path)
        return path

    def test_отчёт_содержит_предпросмотр(self, tmp_path):
        report = check_docx_document(self._doc_with_issues(tmp_path), "issues.docx")
        assert "document_preview" in report
        assert report["document_preview"]["blocks"]

    def test_замечания_привязаны_к_абзацам(self, tmp_path):
        report = check_docx_document(self._doc_with_issues(tmp_path), "issues.docx")
        anchored = [i for i in report["issues"] if "para_index" in i]
        assert anchored, "ни одно замечание не привязано к абзацу"

    def test_привязка_указывает_на_существующий_абзац(self, tmp_path):
        report = check_docx_document(self._doc_with_issues(tmp_path), "issues.docx")
        indices = {b["index"] for b in _paragraphs(report["document_preview"])}
        for issue in report["issues"]:
            if "para_index" in issue:
                assert issue["para_index"] in indices, (
                    f"{issue['code']} ссылается на абзац {issue['para_index']}, "
                    f"которого нет в предпросмотре"
                )

    def test_привязка_совпадает_с_текстом_из_описания(self, tmp_path):
        """location содержит начало абзаца — оно должно совпасть с текстом по индексу."""
        report = check_docx_document(self._doc_with_issues(tmp_path), "issues.docx")
        by_index = {b["index"]: b["text"] for b in _paragraphs(report["document_preview"])}
        checked = 0
        for issue in report["issues"]:
            if "para_index" not in issue or "«" not in issue.get("location", ""):
                continue
            quoted = issue["location"].split("«", 1)[1].rsplit("»", 1)[0].rstrip("…")
            assert by_index[issue["para_index"]].startswith(quoted[:30])
            checked += 1
        assert checked, "не нашлось замечаний с цитатой абзаца"

    @pytest.mark.skipif(
        not (Path(__file__).resolve().parent.parent / "sample_files" / "demo_document_full.docx").exists(),
        reason="нет демо-документа",
    )
    def test_на_демо_документе(self):
        demo = Path(__file__).resolve().parent.parent / "sample_files" / "demo_document_full.docx"
        report = check_docx_document(demo, "demo.docx")
        preview = report["document_preview"]
        assert len(preview["blocks"]) > 5
        assert any(b["kind"] == "table" for b in preview["blocks"])
        assert any("para_index" in i for i in report["issues"])
